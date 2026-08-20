"""Export the raw/mean classifier performance tables to a formatted .xlsx workbook.

Reads rerun/logs/confusion_matrices_mean_raw.json (written by
confusion_matrices_mean_raw.py) and writes a five-sheet workbook intended to be
handed to someone outside this repository: every column is spelled out in full,
every regime is named rather than abbreviated, and each sheet carries its own
explanatory notes so it can be read without this codebase for context.

Sheets:
  Overview                  - what the workbook is, definitions, provenance
  Mortality Classifier      - counts and rates, averaged over the four seeds
  ICU Classifier            - same
  Effect Decomposition      - splits the accuracy change into its two causes
  Per-Seed Detail           - the unaveraged per-seed confusion matrices

    python rerun/export_confusion_workbook.py
"""
import json
import statistics
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "logs" / "confusion_matrices_mean_raw.json"
DESTINATION = HERE / "logs" / "classifier_performance_mean_raw.xlsx"

# Spelled out for the reader of the workbook, who does not have the code.
REGIME_NAMES = {
    "old": "Original (validation merged into test, threshold 0.500)",
    "new@.5": "Control (held-out test only, threshold 0.500)",
    "new": "Corrected (held-out test only, validation-selected threshold)",
}
REGIME_ORDER = ["old", "new@.5", "new"]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
SECTION_FONT = Font(bold=True, size=11, color="1F3864")
NOTE_FONT = Font(italic=True, size=10, color="595959")
HIGHLIGHT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def summarise(rows, regime):
    """Average one regime's confusion matrix across seeds and derive its rates."""
    cells = {k: statistics.mean([r[regime][k] for r in rows])
             for k in ("tp", "fp", "fn", "tn")}
    true_positive = cells["tp"]
    false_positive = cells["fp"]
    false_negative = cells["fn"]
    true_negative = cells["tn"]

    actual_positive = true_positive + false_negative
    actual_negative = false_positive + true_negative
    total = actual_positive + actual_negative
    predicted_positive = true_positive + false_positive

    return {
        "Total Records Evaluated": total,
        "Actual Positive Cases": actual_positive,
        "Actual Negative Cases": actual_negative,
        "True Positives (Correctly Identified Positive)": true_positive,
        "False Positives (Type I Error)": false_positive,
        "False Negatives (Type II Error)": false_negative,
        "True Negatives (Correctly Identified Negative)": true_negative,
        "Overall Accuracy": (true_positive + true_negative) / total,
        "Type I Error Rate (False Positive Rate)": false_positive / actual_negative,
        "Type II Error Rate (False Negative Rate)": false_negative / actual_positive,
        "Sensitivity (Recall, True Positive Rate)": true_positive / actual_positive,
        "Specificity (True Negative Rate)": true_negative / actual_negative,
        "Precision (Positive Predictive Value)": (
            true_positive / predicted_positive if predicted_positive else float("nan")
        ),
        "Prevalence of Positive Class": actual_positive / total,
    }


COUNT_ROWS = [
    "Total Records Evaluated",
    "Actual Positive Cases",
    "Actual Negative Cases",
    "True Positives (Correctly Identified Positive)",
    "False Positives (Type I Error)",
    "False Negatives (Type II Error)",
    "True Negatives (Correctly Identified Negative)",
]
RATE_ROWS = [
    "Overall Accuracy",
    "Type I Error Rate (False Positive Rate)",
    "Type II Error Rate (False Negative Rate)",
    "Sensitivity (Recall, True Positive Rate)",
    "Specificity (True Negative Rate)",
    "Precision (Positive Predictive Value)",
    "Prevalence of Positive Class",
]


def style_header(sheet, row, last_column):
    for column in range(1, last_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def set_widths(sheet, widths):
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_metric_sheet(workbook, title, label, rows, thresholds):
    sheet = workbook.create_sheet(title)
    summaries = {regime: summarise(rows, regime) for regime in REGIME_ORDER}

    sheet["A1"] = f"{title} — Classifier Performance"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        "MIMIC-III, mean aggregation, unfiltered (raw) arm. All counts are the "
        "arithmetic mean of four independent model seeds and are therefore "
        "reported to one decimal place rather than as whole records."
    )
    sheet["A2"].font = NOTE_FONT
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
    sheet.row_dimensions[2].height = 30

    header_row = 4
    sheet.cell(row=header_row, column=1, value="Measure")
    for offset, regime in enumerate(REGIME_ORDER, start=2):
        sheet.cell(row=header_row, column=offset, value=REGIME_NAMES[regime])
    style_header(sheet, header_row, 4)
    sheet.row_dimensions[header_row].height = 46

    current = header_row + 1
    sheet.cell(row=current, column=1, value="Confusion Matrix Counts").font = SECTION_FONT
    current += 1

    for name in COUNT_ROWS:
        sheet.cell(row=current, column=1, value=name).border = BORDER
        for offset, regime in enumerate(REGIME_ORDER, start=2):
            cell = sheet.cell(row=current, column=offset, value=summaries[regime][name])
            cell.number_format = "#,##0.0"
            cell.border = BORDER
            if "Type I" in name or "Type II" in name:
                cell.fill = HIGHLIGHT_FILL
        current += 1

    current += 1
    sheet.cell(row=current, column=1, value="Derived Rates").font = SECTION_FONT
    current += 1

    for name in RATE_ROWS:
        sheet.cell(row=current, column=1, value=name).border = BORDER
        for offset, regime in enumerate(REGIME_ORDER, start=2):
            cell = sheet.cell(row=current, column=offset, value=summaries[regime][name])
            cell.number_format = "0.0000"
            cell.border = BORDER
            if "Type I" in name or "Type II" in name:
                cell.fill = HIGHLIGHT_FILL
        current += 1

    current += 1
    sheet.cell(row=current, column=1, value="Decision Threshold Applied").font = SECTION_FONT
    current += 1
    sheet.cell(row=current, column=1, value="Decision Threshold (mean across seeds)").border = BORDER
    for offset, regime in enumerate(REGIME_ORDER, start=2):
        value = 0.5 if regime in ("old", "new@.5") else statistics.mean(thresholds)
        cell = sheet.cell(row=current, column=offset, value=value)
        cell.number_format = "0.0000"
        cell.border = BORDER
    current += 2

    sheet.cell(row=current, column=1, value="Notes").font = SECTION_FONT
    current += 1
    for note in [
        "A Type I error is a record predicted positive that is actually negative (a false alarm).",
        "A Type II error is a record predicted negative that is actually positive (a missed case).",
        "The Control column shares its test set with the Corrected column and its threshold with "
        "the Original column, so differencing across the three isolates the effect of the smaller "
        "test set from the effect of the moved decision threshold.",
        "The Original column reflects a defect, corrected on 10 August 2026, in which the validation "
        "split was concatenated onto the test split; its test set is therefore approximately twice "
        "the size of the other two columns.",
    ]:
        sheet.cell(row=current, column=1, value=note).font = NOTE_FONT
        current += 1

    set_widths(sheet, [48, 34, 34, 34])
    sheet.freeze_panes = "B5"
    return summaries


def write_overview(workbook):
    sheet = workbook.create_sheet("Overview", 0)
    sheet["A1"] = "Classifier Performance — Validation Split Correction"
    sheet["A1"].font = TITLE_FONT

    content = [
        ("", ""),
        ("Purpose", "Quantifies how the 10 August 2026 correction to the train/validation/test "
                    "split changed measured classifier performance, and separates that change "
                    "into its two independent causes."),
        ("Dataset", "MIMIC-III intensive care unit records, 46,032 stays, seven bedside vitals "
                    "over a 24-hour window."),
        ("Aggregation Method", "Mean"),
        ("Filter Arm", "Raw (unfiltered baseline)"),
        ("Classifier", "Random forest, 300 trees, scikit-learn implementation"),
        ("Seeds", "22, 985, 439, 81 — reported figures are the mean across all four"),
        ("Prediction Labels", "Mortality and intensive care unit admission, evaluated separately"),
        ("", ""),
        ("Definition — Type I Error", "A record predicted positive that is actually negative. "
                                      "Also called a false positive or false alarm."),
        ("Definition — Type II Error", "A record predicted negative that is actually positive. "
                                       "Also called a false negative or missed case."),
        ("Definition — Sensitivity", "Proportion of actually positive records correctly identified."),
        ("Definition — Specificity", "Proportion of actually negative records correctly identified."),
        ("Definition — Precision", "Proportion of records predicted positive that are actually positive."),
        ("", ""),
        ("Regime — Original", "The pre-correction configuration. The validation split was "
                              "concatenated onto the test split, and predictions were taken at the "
                              "implicit 0.500 decision threshold."),
        ("Regime — Control", "Added for this analysis. Uses the corrected held-out test split but "
                             "retains the 0.500 threshold, isolating the effect of test set size."),
        ("Regime — Corrected", "The current configuration. Test is a true held-out ten percent and "
                               "the decision threshold is selected on the validation split by "
                               "maximising Youden's J statistic."),
        ("", ""),
        ("Provenance", "Computed by rerun/confusion_matrices_mean_raw.py under SLURM job 19568266 "
                       "on 11 August 2026. Source data: rerun/logs/confusion_matrices_mean_raw.json."),
        ("Verification", "Reproduces the decision thresholds and accuracy figures stored in the "
                         "filter impact result files exactly."),
    ]

    row = 2
    for heading, body in content:
        if heading:
            cell = sheet.cell(row=row, column=1, value=heading)
            cell.font = SECTION_FONT
            cell.alignment = Alignment(vertical="top")
            body_cell = sheet.cell(row=row, column=2, value=body)
            body_cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    set_widths(sheet, [30, 95])
    return sheet


def write_decomposition(workbook, per_label_summaries):
    sheet = workbook.create_sheet("Effect Decomposition")
    sheet["A1"] = "Decomposition of the Change in Overall Accuracy"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        "The total change is split into the part attributable to the smaller held-out "
        "test set and the part attributable to the moved decision threshold."
    )
    sheet["A2"].font = NOTE_FONT
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)

    headers = [
        "Prediction Label",
        "Effect",
        "Change in Overall Accuracy",
        "Share of Total Change",
        "Description",
    ]
    for column, name in enumerate(headers, start=1):
        sheet.cell(row=4, column=column, value=name)
    style_header(sheet, 4, len(headers))
    sheet.row_dimensions[4].height = 32

    row = 5
    for label, summaries in per_label_summaries.items():
        original = summaries["old"]["Overall Accuracy"]
        control = summaries["new@.5"]["Overall Accuracy"]
        corrected = summaries["new"]["Overall Accuracy"]

        size_effect = control - original
        threshold_effect = corrected - control
        total = corrected - original

        entries = [
            ("Reduced Test Set Size", size_effect,
             "Effect of no longer merging the validation split into the test split."),
            ("Moved Decision Threshold", threshold_effect,
             "Effect of selecting the threshold on validation rather than using 0.500."),
            ("Total Change", total, "Sum of the two effects above."),
        ]
        for name, value, description in entries:
            sheet.cell(row=row, column=1, value=label).border = BORDER
            sheet.cell(row=row, column=2, value=name).border = BORDER
            accuracy_cell = sheet.cell(row=row, column=3, value=value)
            accuracy_cell.number_format = "+0.0000;-0.0000"
            accuracy_cell.border = BORDER
            share_cell = sheet.cell(
                row=row, column=4,
                value=(value / total if name != "Total Change" else 1.0),
            )
            share_cell.number_format = "0.0%"
            share_cell.border = BORDER
            description_cell = sheet.cell(row=row, column=5, value=description)
            description_cell.border = BORDER
            description_cell.alignment = Alignment(wrap_text=True, vertical="top")
            if name == "Total Change":
                for column in range(1, 6):
                    sheet.cell(row=row, column=column).font = Font(bold=True)
            row += 1
        row += 1

    set_widths(sheet, [20, 30, 26, 20, 62])
    sheet.freeze_panes = "A5"


def write_per_seed(workbook, data):
    sheet = workbook.create_sheet("Per-Seed Detail")
    sheet["A1"] = "Per-Seed Confusion Matrices (Unaveraged)"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        "The underlying figures behind the averaged tables. Note that all four seeds "
        "share a single data partition — only the random forest's own randomness "
        "varies — so the spread across seeds understates true sampling variability."
    )
    sheet["A2"].font = NOTE_FONT
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
    sheet.row_dimensions[2].height = 30

    headers = [
        "Prediction Label", "Seed", "Evaluation Regime", "Decision Threshold",
        "True Positives", "False Positives (Type I Error)",
        "False Negatives (Type II Error)", "True Negatives",
    ]
    for column, name in enumerate(headers, start=1):
        sheet.cell(row=4, column=column, value=name)
    style_header(sheet, 4, len(headers))
    sheet.row_dimensions[4].height = 44

    row = 5
    for label in ("mortality", "icu"):
        display = "Mortality" if label == "mortality" else "Intensive Care Unit"
        for entry in data[label]:
            for regime in REGIME_ORDER:
                matrix = entry[regime]
                threshold = 0.5 if regime in ("old", "new@.5") else entry["threshold"]
                values = [
                    display, entry["seed"], REGIME_NAMES[regime], threshold,
                    matrix["tp"], matrix["fp"], matrix["fn"], matrix["tn"],
                ]
                for column, value in enumerate(values, start=1):
                    cell = sheet.cell(row=row, column=column, value=value)
                    cell.border = BORDER
                    if column == 4:
                        cell.number_format = "0.0000"
                    elif column >= 5:
                        cell.number_format = "#,##0"
                row += 1

    set_widths(sheet, [22, 10, 46, 20, 16, 26, 26, 16])
    sheet.freeze_panes = "A5"


def main():
    data = json.loads(SOURCE.read_text())

    workbook = Workbook()
    workbook.remove(workbook.active)

    write_overview(workbook)

    summaries = {}
    summaries["Mortality"] = write_metric_sheet(
        workbook, "Mortality Classifier", "mortality", data["mortality"],
        [entry["threshold"] for entry in data["mortality"]],
    )
    summaries["Intensive Care Unit"] = write_metric_sheet(
        workbook, "ICU Classifier", "icu", data["icu"],
        [entry["threshold"] for entry in data["icu"]],
    )

    write_decomposition(workbook, summaries)
    write_per_seed(workbook, data)

    workbook.save(DESTINATION)
    print(f"wrote {DESTINATION}")


if __name__ == "__main__":
    main()
